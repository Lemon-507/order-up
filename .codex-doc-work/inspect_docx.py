from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def font_name(run):
    rpr = run._element.rPr
    if rpr is None or rpr.rFonts is None:
        return run.font.name
    return (
        rpr.rFonts.get(qn("w:eastAsia"))
        or rpr.rFonts.get(qn("w:ascii"))
        or rpr.rFonts.get(qn("w:hAnsi"))
        or run.font.name
    )


def para_record(p, index, part="document"):
    return {
        "part": part,
        "index": index,
        "style": p.style.name if p.style else None,
        "text": p.text,
        "alignment": str(p.alignment),
        "runs": [
            {
                "text": r.text,
                "font": font_name(r),
                "size_pt": r.font.size.pt if r.font.size else None,
                "bold": r.bold,
                "italic": r.italic,
            }
            for r in p.runs
            if r.text
        ],
    }


def main():
    source = Path(sys.argv[1])
    out = Path(sys.argv[2])
    doc = Document(source)
    data = {
        "source": str(source),
        "paragraphs": [para_record(p, i) for i, p in enumerate(doc.paragraphs)],
        "tables": [],
        "sections": [],
        "headers": [],
        "footers": [],
        "parts": [],
    }
    for ti, table in enumerate(doc.tables):
        data["tables"].append({
            "index": ti,
            "style": table.style.name if table.style else None,
            "rows": [
                [
                    {
                        "text": cell.text,
                        "paragraphs": [para_record(p, pi, f"table[{ti}]") for pi, p in enumerate(cell.paragraphs)],
                    }
                    for cell in row.cells
                ]
                for row in table.rows
            ],
        })
    for si, section in enumerate(doc.sections):
        data["sections"].append({
            "index": si,
            "page_width": section.page_width,
            "page_height": section.page_height,
            "top_margin": section.top_margin,
            "bottom_margin": section.bottom_margin,
            "left_margin": section.left_margin,
            "right_margin": section.right_margin,
            "header_distance": section.header_distance,
            "footer_distance": section.footer_distance,
            "start_type": str(section.start_type),
        })
        data["headers"].append([para_record(p, i, f"header[{si}]") for i, p in enumerate(section.header.paragraphs)])
        data["footers"].append([para_record(p, i, f"footer[{si}]") for i, p in enumerate(section.footer.paragraphs)])
    with zipfile.ZipFile(source) as zf:
        data["parts"] = [
            {"path": info.filename, "size": info.file_size, "crc": info.CRC}
            for info in zf.infolist()
        ]
    out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
