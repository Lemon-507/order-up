from __future__ import annotations

import hashlib
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


REFERENCE = Path(r"C:\Users\shibo\Desktop\JavaFX-doc\Day04(9.10)\XX项目_详细设计说明书_v3.0.docx")
WORK = Path(r"C:\Users\shibo\Desktop\order-up\.codex-doc-work")
OUTPUT = Path(r"C:\Users\shibo\Desktop\order-up\出餐！出餐！_详细设计说明书_v3.0.docx")
EXPECTED_SHA256 = "b0a36132542947a9fef1875aa82d19309d671567c691728c7451fc31418a44e0"

CN_FONT = r"C:\Windows\Fonts\msyh.ttc"
CN_BOLD_FONT = r"C:\Windows\Fonts\msyhbd.ttc"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def set_run_font(run, name="宋体", size=12, bold=None, italic=None, color="000000"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=90, bottom=85, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="808080", size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), str(size))
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if value and node is None:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)
    elif not value and node is not None:
        p_pr.remove(node)


def set_cell_text(cell, text, *, bold=False, align="left", font="宋体", size=9.2, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    run = p.add_run(str(text))
    set_run_font(run, font, size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths_cm=None, aligns=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table"
    table.autofit = False
    table.alignment = 1
    set_table_borders(table, color="A6A6A6", size=5)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, align="center", size=9.5, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[i], "404040")
        if widths_cm:
            table.rows[0].cells[i].width = Cm(widths_cm[i])
    for ri, row_values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_values):
            align = aligns[i] if aligns else "left"
            set_cell_text(row.cells[i], value, align=align, size=font_size)
            if widths_cm:
                row.cells[i].width = Cm(widths_cm[i])
            if ri % 2 == 1:
                set_cell_shading(row.cells[i], "F2F2F2")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(13 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_run_font(r, "黑体", 14 if level == 1 else 12, bold=True)
    return p


def add_body(doc, text="", *, bold_lead=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(style="Normal")
    p.alignment = align
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, "宋体", 10.5, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, "宋体", 10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, "宋体", 10.5)
    return p


def add_note(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.right_indent = Pt(18)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_run_font(r, "宋体", 9.5, color="595959")
    return p


def add_code(doc, code):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.right_indent = Pt(12)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code)
    set_run_font(r, "Consolas", 8.8)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    p._p.get_or_add_pPr().append(shading)
    return p


def arrow(draw, start, end, color=(65, 65, 65), width=4, dashed=False):
    x1, y1 = start
    x2, y2 = end
    if dashed:
        segments = 12
        for i in range(segments):
            if i % 2 == 0:
                a = i / segments
                b = (i + 1) / segments
                draw.line((x1 + (x2-x1)*a, y1 + (y2-y1)*a, x1 + (x2-x1)*b, y1 + (y2-y1)*b), fill=color, width=width)
    else:
        draw.line((x1, y1, x2, y2), fill=color, width=width)
    angle = math.atan2(y2-y1, x2-x1)
    l = 15
    for delta in (2.55, -2.55):
        draw.line((x2, y2, x2 + l*math.cos(angle+delta), y2 + l*math.sin(angle+delta)), fill=color, width=width)


def center_text(draw, box, text, font, fill=(0,0,0)):
    x1, y1, x2, y2 = box
    lines = text.split("\n")
    heights = []
    widths = []
    for line in lines:
        bb = draw.textbbox((0,0), line, font=font)
        widths.append(bb[2]-bb[0])
        heights.append(bb[3]-bb[1])
    total = sum(heights) + 8*(len(lines)-1)
    y = y1 + (y2-y1-total)/2
    for line, w, h in zip(lines, widths, heights):
        draw.text((x1+(x2-x1-w)/2, y), line, font=font, fill=fill)
        y += h + 8


def rounded_node(draw, box, title, subtitle, fill, outline, title_font, body_font):
    draw.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=3)
    x1,y1,x2,y2=box
    center_text(draw,(x1,y1+8,x2,y1+46),title,title_font)
    center_text(draw,(x1+8,y1+44,x2-8,y2-8),subtitle,body_font,fill=(70,70,70))


def make_main_state_diagram(path: Path):
    img = Image.new("RGB", (1200, 470), "white")
    d = ImageDraw.Draw(img)
    title = ImageFont.truetype(CN_BOLD_FONT, 30)
    body = ImageFont.truetype(CN_FONT, 20)
    label = ImageFont.truetype(CN_FONT, 19)
    ready=(60,95,270,220); running=(410,95,650,220); finished=(860,95,1100,220); paused=(410,315,650,430)
    rounded_node(d,ready,"READY","枚举已定义\n无独立状态字段",(247,247,247),(90,90,90),title,body)
    rounded_node(d,running,"RUNNING","startGame 后\n倒计时与帧更新运行",(232,242,252),(31,78,121),title,body)
    rounded_node(d,finished,"FINISHED","倒计时到零或\nstopGame 结束",(242,242,242),(90,90,90),title,body)
    rounded_node(d,paused,"PAUSED","待实现\n无暂停入口与恢复方法",(255,246,224),(160,110,20),title,body)
    arrow(d,(270,157),(410,157)); center_text(d,(275,108,405,145),"startGame",label)
    arrow(d,(650,157),(860,157)); center_text(d,(655,98,855,145),"timeUp / stop",label)
    arrow(d,(530,220),(530,315),dashed=True); center_text(d,(540,235,770,280),"暂停入口待实现",label)
    center_text(d,(50,10,1150,60),"当前源码可验证的游戏生命周期",title)
    img.save(path, dpi=(180,180))


def make_sub_state_diagram(path: Path):
    img = Image.new("RGB", (1100, 370), "white")
    d = ImageDraw.Draw(img)
    title = ImageFont.truetype(CN_BOLD_FONT, 30)
    body = ImageFont.truetype(CN_FONT, 20)
    label = ImageFont.truetype(CN_FONT, 19)
    idle=(130,120,390,255); moving=(710,120,970,255)
    rounded_node(d,idle,"IDLE","pressedDirections 为空\n位置保持",(247,247,247),(90,90,90),title,body)
    rounded_node(d,moving,"MOVING","方向集合非空\n按 deltaSeconds 更新位置",(232,242,252),(31,78,121),title,body)
    arrow(d,(390,155),(710,155)); center_text(d,(400,95,700,140),"press(W/A/S/D)",label)
    arrow(d,(710,225),(390,225)); center_text(d,(405,235,700,285),"release 最后一个方向",label)
    center_text(d,(50,20,1050,70),"RUNNING 状态下的玩家移动输入子状态",title)
    img.save(path, dpi=(180,180))


def add_picture(doc, path, width_cm, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    picture = p.add_run().add_picture(str(path), width=Cm(width_cm))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run_font(r, "宋体", 9.5)


def clear_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def add_cover(doc):
    for _ in range(7):
        doc.add_paragraph()
    for text, size in (("出餐 出餐", 24), ("详细设计说明书", 24)):
        p = doc.add_paragraph(style="主标题")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        r = p.add_run(text)
        set_run_font(r, "宋体", size, bold=True)
    for _ in range(8):
        doc.add_paragraph()
    metadata = [
        "小组成员：待确认（项目仓库未提供成员名单）",
        "2026年9月10日",
        "版本：V3.0",
    ]
    for text in metadata:
        p = doc.add_paragraph(style="主标题")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        r = p.add_run(text)
        set_run_font(r, "宋体", 15, bold=True)
    doc.add_page_break()


def method_rows(class_name):
    data = {
        "Launcher": [
            ("start", "void start(Stage stage)", "stage：主窗口", "无", "JavaFX 调用 Application.start", "设置窗口标题、最小尺寸并显示开始页"),
            ("showScene", "void showScene(String fxmlPath)", "fxmlPath：类路径资源", "无", "资源存在且 FXML 可加载", "释放旧 GameView，切换根节点并配置控制器"),
            ("configureController", "void configureController(Object controller)", "controller：FXML 控制器", "无", "FXML 已加载", "为开始页、结果页或游戏页注入页面切换回调"),
            ("playMenuMusic", "void playMenuMusic()", "无", "无", "startmenu.mp3 可选存在", "创建或复用 MediaPlayer 并循环播放，音量 0.35"),
            ("stop", "void stop()", "无", "无", "应用退出", "释放 GameView 与菜单 MediaPlayer"),
        ],
        "GameView": [
            ("initialize", "void initialize()", "无", "无", "FXML 字段已注入", "创建 GameController，启动 60 秒倒计时、输入与游戏循环"),
            ("startGameLoop", "void startGameLoop()", "无", "无", "controller 与 Canvas 已初始化", "每帧计算不超过 0.05 秒的 deltaSeconds 并更新、绘制"),
            ("renderFrame", "void renderFrame(GraphicsContext graphics)", "graphics：画布绘图上下文", "无", "graphics 非空", "清空画布，绘制玩家并刷新倒计时显示"),
            ("renderTime", "void renderTime(int totalSeconds)", "totalSeconds：剩余秒数", "无", "秒数由 GameTimer 提供", "显示 mm:ss；剩余 10 秒及以下变为红色"),
            ("dispose", "void dispose()", "无", "无", "页面将被替换或应用退出", "停止 AnimationTimer 和 GameController"),
        ],
        "GameController": [
            ("构造方法", "GameController(double worldWidth, double worldHeight, int gameSeconds, Runnable onGameFinished)", "世界尺寸、局时长、结束回调", "新实例", "参数由 GameView 提供", "玩家初始位置为 200,200，并创建 GameTimer"),
            ("startGame", "void startGame()", "无", "无", "gameSeconds 大于 0", "finished=false，启动倒计时"),
            ("update", "void update(double deltaSeconds)", "deltaSeconds：本帧秒数", "无", "游戏未结束；deltaSeconds 非负", "更新玩家位置和倒计时"),
            ("finishGame", "void finishGame()", "无", "无", "尚未结束", "幂等地停止计时、清空输入并调用结束回调"),
            ("stopGame", "void stopGame()", "无", "无", "页面销毁", "直接标记结束并停止计时、清空输入"),
        ],
        "Player": [
            ("press", "void press(Direction direction)", "direction：方向枚举", "无", "direction 非空", "将方向加入 EnumSet"),
            ("release", "void release(Direction direction)", "direction：方向枚举", "无", "无", "从 EnumSet 移除方向"),
            ("clearMovement", "void clearMovement()", "无", "无", "无", "清空全部已按下方向"),
            ("update", "void update(double deltaSeconds, double worldWidth, double worldHeight)", "帧时间与世界尺寸", "无", "世界尺寸不小于玩家尺寸", "按速度更新位置，并限制在画布边界内"),
            ("setSpeed", "void setSpeed(double speed)", "speed：像素每秒", "无", "源码未校验非负值", "更新移动速度；负值处理规则待确认"),
        ],
        "GameTimer": [
            ("构造方法", "GameTimer(Runnable onTimeUp)", "onTimeUp：超时回调", "新实例", "回调应非空；源码未校验", "保存回调，计时尚未运行"),
            ("startCountDown", "void startCountDown(int totalSeconds)", "totalSeconds：总秒数", "无", "totalSeconds > 0", "重置累计秒数并开始倒计时"),
            ("update", "void update(double deltaSeconds)", "deltaSeconds：本帧秒数", "无", "运行中且 deltaSeconds >= 0", "累计满 1 秒则递减，归零时仅回调一次"),
            ("stop", "void stop()", "无", "无", "无", "running=false，保留剩余秒数"),
            ("getSecondsCount", "int getSecondsCount()", "无", "剩余秒数", "无", "不修改状态"),
        ],
    }
    return data[class_name]


def build():
    if sha256(REFERENCE) != EXPECTED_SHA256:
        raise RuntimeError("Reference template hash changed; distillation must be repeated.")
    WORK.mkdir(parents=True, exist_ok=True)
    main_diagram = WORK / "main-state.png"
    sub_diagram = WORK / "movement-substate.png"
    make_main_state_diagram(main_diagram)
    make_sub_state_diagram(sub_diagram)

    doc = Document(REFERENCE)
    clear_body(doc)
    styles = doc.styles
    for style_name in ("主标题", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.color.rgb = RGBColor(0, 0, 0)
    add_cover(doc)

    add_heading(doc, "文档范围", 1)
    add_body(doc, "本文档描述仓库当前版本中可由源码、资源文件和 Git 历史验证的详细设计。项目采用 Java 17、JavaFX 17.0.10、FXML、Canvas 和 Maven；当前已形成页面切换、菜单音乐、60 秒倒计时及玩家移动的最小运行闭环。")
    add_body(doc, "实现状态分为已实现、部分实现、未实现和待确认。README 或需求分析中的目标只有在源码存在可执行落点时才记为已实现；只有类型、接口或空方法时记为部分实现或未实现。")
    add_note(doc, "待确认：小组成员名单、正式验收口径、关卡数量和菜谱内容未出现在当前工作树中。AI 使用信息依据 Git 历史中的 AI记录.md；没有逐方法标注的内容保留为不确定。")

    add_heading(doc, "一 状态机图与状态转移表", 1)
    add_heading(doc, "1.1 主状态机", 2)
    add_body(doc, "GameState 定义 READY、RUNNING、PAUSED 和 FINISHED 四个状态，但当前 GameController 没有 currentState 字段，也没有统一的 setState 方法。实际生命周期由 finished 布尔值、GameTimer.running 和页面回调共同表达。下表把枚举中的设计意图与当前可执行路径分开列出。")
    add_table(doc,
        ["当前状态", "触发事件", "前置条件", "下一状态", "说明"],
        [
            ("READY", "GameView.initialize 调用 startGame", "FXML 已加载，gameSeconds=60", "RUNNING", "实际代码启动计时；未写入 GameState 字段"),
            ("RUNNING", "AnimationTimer 每帧回调", "finished=false", "RUNNING", "更新 Player 和 GameTimer"),
            ("RUNNING", "倒计时归零", "secondsCount 递减至 0", "FINISHED", "GameTimer 调用 finishGame，随后切换结果页"),
            ("RUNNING", "GameView.dispose 调用 stopGame", "游戏页面被替换或应用退出", "FINISHED", "停止循环和计时；不触发结果回调"),
            ("RUNNING", "暂停操作", "待确认：暂停入口与按键", "PAUSED", "未实现"),
            ("PAUSED", "恢复操作", "待确认：恢复入口", "RUNNING", "未实现"),
            ("FINISHED", "结果页 restartGame", "结果页控制器已配置", "READY 后进入 RUNNING", "回到 game.fxml 后新建 GameController；result.fxml 当前为空，按钮未落地"),
        ],
        [2.0, 4.0, 4.0, 2.1, 5.0], ["center", "left", "left", "center", "left"], 8.8)

    add_heading(doc, "1.2 状态枚举定义", 2)
    add_code(doc, "public enum GameState {\n    READY,\n    RUNNING,\n    PAUSED,\n    FINISHED\n}")
    add_note(doc, "该枚举存在于 com.orderup.model.GameState，但当前无其他源码引用。PAUSED 因此仅是预留设计，不能视为已实现功能。")

    add_heading(doc, "1.3 状态切换方法", 2)
    add_body(doc, "当前没有模板示例中的 setState 方法。可执行切换分别落在 GameController.startGame、finishGame 和 stopGame 中。finishGame 通过 finished 标志保证回调只执行一次。")
    add_code(doc, "public void startGame() {\n    finished = false;\n    gameTimer.startCountDown(gameSeconds);\n}\n\npublic void finishGame() {\n    if (finished) return;\n    finished = true;\n    gameTimer.stop();\n    player.clearMovement();\n    onGameFinished.run();\n}")

    add_heading(doc, "1.4 主状态机图", 2)
    add_picture(doc, main_diagram, 15.8, "图 1 当前源码可验证的主状态流转")

    add_heading(doc, "1.5 状态机自查", 2)
    add_table(doc,
        ["检查项", "结果", "说明"],
        [
            ("无死状态", "部分满足", "FINISHED 可经重新加载游戏页重新开始；PAUSED 无进入和退出实现"),
            ("无不可达状态", "不满足", "GameState.PAUSED 在当前源码中不可达；GameState 整体未接入控制器"),
            ("覆盖 P0 全部流程", "不满足", "当前仅覆盖启动、计时、移动和结束；订单、加工、装盘、出餐、评分、暂停等未形成闭环"),
        ], [5.5, 3.0, 8.6], ["left", "center", "left"], 9.0)

    add_heading(doc, "1.6 子状态机", 2)
    add_body(doc, "玩家移动输入可视为 RUNNING 下的概念子状态。源码没有为它定义枚举，而是使用 EnumSet<Direction> pressedDirections：集合为空时玩家静止，集合非空时按各方向合成位移。")
    add_table(doc,
        ["当前子状态", "触发事件", "前置条件", "下一子状态", "说明"],
        [
            ("IDLE", "按下 W A S D 任一键", "Canvas 获得焦点", "MOVING", "GameView 转换为 Direction 并调用 controller.press"),
            ("MOVING", "继续按住或新增方向键", "游戏未结束", "MOVING", "允许多方向同时存在；对角线速度未归一化"),
            ("MOVING", "释放部分方向键", "仍有其他方向键按下", "MOVING", "移除对应 Direction"),
            ("MOVING", "释放最后一个方向键", "方向集合变为空", "IDLE", "位置停止更新"),
            ("MOVING", "Canvas 失焦或游戏结束", "监听到失焦或调用 finishGame", "IDLE", "clearMovement 清空输入"),
        ], [2.5, 4.2, 3.8, 2.5, 4.6], ["center", "left", "left", "center", "left"], 8.7)

    add_heading(doc, "1.7 子状态机图", 2)
    add_picture(doc, sub_diagram, 15.5, "图 2 玩家移动输入子状态")

    add_heading(doc, "二 方法级设计", 1)
    add_body(doc, "以下方法签名和行为均来自当前源码。对空值、负速度等未校验情形只描述现状，不补写未存在的约束。")
    responsibilities = {
        "Launcher": "JavaFX 应用入口、窗口配置、FXML 页面切换和菜单音乐生命周期",
        "GameView": "FXML 显示层、键盘输入、逐帧循环、画布绘制与倒计时展示",
        "GameController": "协调玩家、计时器和游戏结束条件",
        "Player": "保存移动输入、位置和速度并执行边界限制",
        "GameTimer": "与 JavaFX 控件解耦的帧驱动倒计时",
    }
    for idx, class_name in enumerate(responsibilities, 1):
        add_heading(doc, f"2.{idx} {class_name}", 2)
        add_body(doc, f"职责：{responsibilities[class_name]}。")
        add_table(doc,
            ["方法名", "方法签名", "参数说明", "返回值", "前置条件", "后置条件"],
            method_rows(class_name),
            [2.0, 5.0, 3.4, 1.7, 3.5, 4.1], ["center", "left", "left", "center", "left", "left"], 8.2)

    add_heading(doc, "2.6 其余模型与服务骨架", 2)
    add_table(doc,
        ["类别", "主要类型", "当前设计落点", "实现状态"],
        [
            ("订单", "Order Recipe OrderResult", "订单字段、菜谱需求、结果字段", "数据模型已定义；字段类型仍较粗，如 remainingTime 和 success 使用 String"),
            ("食材", "Ingredient IngredientType IngredientStatus", "食材类型、RAW CUT COOKED 状态和进度", "仅 FISH 一种类型，未接入游戏循环"),
            ("工作台", "Station StationType", "食材库、切菜板、煎锅、煮锅、搅拌机、出餐口和垃圾桶枚举", "模型存在，交互逻辑未实现"),
            ("装盘", "Plate", "Set<Ingredient> contents", "模型存在，未与订单提交集成"),
            ("订单服务", "OrderService OrderServiceImpl", "创建、查询、更新和提交接口", "实现方法为 null、List.of 或空方法"),
            ("评分", "ScoreService CalculateScore", "成功分、小费、罚分接口；星级阈值计算", "ScoreService 无实现；CalculateScore 未接入游戏"),
            ("其他服务", "GameService PlayerService KitchenService", "分层接口占位", "接口为空；仅 GameService 和 PlayerService 有空实现类"),
            ("关卡配置", "GameConfig GameMap", "关卡与 13×9 地图尺寸占位", "GameConfig 为空；GameMap 仅静态行列字段"),
        ], [2.2, 4.2, 6.7, 5.0], ["center", "left", "left", "left"], 8.6)

    doc.add_page_break()
    add_heading(doc, "三 P0 设计落点对照", 1)
    add_body(doc, "P0 功能点来自仓库中的需求分析.md。设计落点按当前源码判断，不把未接入的枚举、接口或 README 描述记为完成。")
    p0_rows = [
        ("暂停、重新开始、退出", "StartController.onQuitButtonClick；ResultController.restartGame；GameState.PAUSED", "部分实现", "退出与重启回调存在；暂停入口、暂停计时和恢复流程未实现"),
        ("玩家全向移动、拿取物品、操作设备", "GameView 键盘映射；GameController.press/release；Player.update", "部分实现", "WASD 移动和边界限制已实现；拿取与设备交互未实现"),
        ("食材库无限拿取并区分种类", "StationType.INGREDIENT_SOURCE；IngredientType", "未实现", "仅有类型占位；IngredientType 当前只有 FISH，无拿取服务"),
        ("顾客限时持续提出订单", "Order Recipe；OrderService.createRandomOrder/updateOrders", "未实现", "OrderServiceImpl 返回 null、空列表或空操作"),
        ("完成订单后评分和时效小费", "ScoreService；CalculateScore", "未实现", "评分接口与星级算法存在，但无 ScoreService 实现且未接入订单"),
        ("初始页、关卡选择页、两个以上游戏场景", "start.fxml game.fxml result.fxml；Launcher", "部分实现", "三类页面资源存在；关卡选择页和两个不同游戏地图未实现，result.fxml 为空"),
        ("游戏计时器", "GameTimer；GameView.renderTime", "已实现", "60 秒倒计时、mm:ss 展示、最后 10 秒红色提示和归零结束已接通"),
        ("角色移动动画", "AnimationTimer；PlayerView.render", "部分实现", "逐帧刷新已实现；角色目前仅为红色矩形，没有精灵帧动画"),
        ("背景音乐和音效", "Launcher.playMenuMusic；startmenu.mp3", "部分实现", "开始菜单音乐可循环播放；仓库未见游戏过程音效资源和调用"),
        ("煎锅、煮锅、切菜台、搅拌机", "StationType 枚举", "未实现", "类型名称已定义，设备实体、加工进度和交互服务未实现"),
        ("出餐口", "StationType.SERVING_COUNTER；OrderService.submitPlate", "未实现", "类型与接口存在，提交实现返回 null"),
        ("盘子组合食材形成菜肴", "Plate.contents；Recipe.requirements", "未实现", "数据结构存在，匹配、组合和转移规则未实现"),
        ("垃圾桶丢弃食材和菜肴", "StationType.TRASH_BIN", "未实现", "只有枚举值，无清空物品交互"),
    ]
    add_table(doc, ["P0 功能点", "设计落点", "状态", "设计说明"], p0_rows,
              [4.2, 5.2, 2.0, 6.8], ["left", "left", "center", "left"], 8.1)

    add_heading(doc, "3.1 设计模式使用说明", 2)
    add_table(doc,
        ["设计决策", "使用位置", "选择理由或当前限制"],
        [
            ("MVC 式分层", "model controller view service 包", "显示、状态协调、数据模型和服务接口分开；但 GameView 同时处理输入和绘制，属于轻量实现"),
            ("回调注入", "Launcher.configureController；GameController.onGameFinished", "用 Runnable 解耦页面控制器与 Stage；没有引入额外事件总线"),
            ("帧时间驱动", "GameView.startGameLoop；Player.update；GameTimer.update", "移动与计时使用 deltaSeconds，不依赖固定帧率；每帧时间上限为 0.05 秒"),
            ("接口加多实现的策略模式", "OrderService ScoreService 等", "当前没有多个有效实现，不能认定已形成策略模式；接口多为骨架"),
            ("观察者式变化通知", "Canvas focusedProperty 监听器", "仅对焦点变化使用 JavaFX 监听器；游戏状态与订单尚未建立统一观察者机制"),
            ("未使用的抽象层", "GameService PlayerService KitchenService 及空实现", "目前没有行为或调用方；应在功能落地时补全，否则后续可删除。是否原计划删除待确认"),
        ], [4.2, 5.1, 8.9], ["left", "left", "left"], 8.5)

    add_heading(doc, "四 AI 使用与核对说明", 1)
    add_heading(doc, "4.1 AI 使用记录", 2)
    add_body(doc, "仓库历史提交 84cd710 中保存了 AI记录.md，可验证以下使用记录。记录未逐条注明日期，也未说明每段示例是否直接复制进源码。")
    add_table(doc,
        ["序号", "AI 工具", "生成内容", "提示词摘要"],
        [
            ("1", "Gemini 3.6 Flash", "中英文项目名称建议、README 模板、倒计时示例", "游戏命名；README 模板；JavaFX 正计时和倒计时设计"),
            ("2", "GPT 5.5", "10 天 4 人开发路线和功能取舍", "检查 README 并规划实训开发路线"),
            ("3", "GPT 5.6", "JavaFX Canvas 使用、绘图和动画示例", "Canvas 如何应用于 JavaFX 项目"),
            ("4", "OpenAI Codex", "本详细设计说明书；未生成项目代码", "依据模板和仓库证据生成详细设计，并标注不确定项"),
        ], [1.4, 3.8, 6.1, 7.0], ["center", "center", "left", "left"], 8.7)

    add_heading(doc, "4.2 AI 生成代码核对", 2)
    add_heading(doc, "4.2.1 倒计时设计", 2)
    add_body(doc, "AI 记录中的原始示例把 Timeline、Label 和按钮事件写在同一个 Application 中，并由 updateTimer 直接更新 UI。")
    add_code(doc, "timeline = new Timeline(new KeyFrame(Duration.seconds(1), e -> updateTimer()));\n\nprivate void updateTimer() {\n    secondsCount--;\n    if (secondsCount <= 0) {\n        timeline.stop();\n        onTimeUp();\n    }\n    renderTime();\n}")
    add_table(doc,
        ["问题类型", "位置", "问题描述", "修正方式"],
        [
            ("职责耦合", "AI 示例 GameTimerApp", "计时状态、JavaFX 控件和按钮事件集中在 Application，难以独立测试或复用", "当前 GameTimer 不依赖 JavaFX；GameView 负责显示，GameController 负责结束回调"),
            ("与游戏循环重复", "Timeline 每秒触发", "项目已有 AnimationTimer；再维护 Timeline 会形成两个时间源", "当前由 GameView 统一传入 deltaSeconds，GameTimer 累计满一秒再递减"),
        ], [2.5, 4.2, 6.2, 6.3], ["center", "left", "left", "left"], 8.3)
    add_body(doc, "当前实现的核心代码如下。它修正了职责耦合和双时间源问题，但没有 pause 或 resume 方法，因此 P0 暂停功能仍未完成。")
    add_code(doc, "public void update(double deltaSeconds) {\n    if (!running) return;\n    if (deltaSeconds < 0) {\n        throw new IllegalArgumentException(\"Delta seconds cannot be negative.\");\n    }\n    accumulatedSeconds += deltaSeconds;\n    while (running && accumulatedSeconds >= 1) {\n        accumulatedSeconds -= 1;\n        secondsCount--;\n        if (secondsCount <= 0) {\n            secondsCount = 0;\n            running = false;\n            onTimeUp.run();\n        }\n    }\n}")

    add_heading(doc, "4.2.2 Canvas 移动设计", 2)
    add_body(doc, "AI 记录中的 Canvas 动画示例按每帧固定增加 2 像素，并把绘制与位置更新写在同一段 handle 逻辑中。")
    add_code(doc, "public void handle(long now) {\n    gc.clearRect(0, 0, canvas.getWidth(), canvas.getHeight());\n    gc.fillOval(x, 100, 50, 50);\n    x += 2;\n}")
    add_table(doc,
        ["问题类型", "位置", "问题描述", "修正方式"],
        [
            ("帧率相关", "AI 示例 x += 2", "不同帧率会产生不同移动速度", "Player.update 使用 speed × deltaSeconds，默认速度 220 像素每秒"),
            ("职责混合", "AI 示例 handle", "位置状态和绘制代码混在 AnimationTimer 内", "当前 Player 保存状态，GameController 更新，PlayerView 绘制，GameView 驱动循环"),
            ("时间突跳", "页面卡顿后的首帧", "过大的时间差可能让玩家或计时器一次跳过过多", "GameView 将单帧 deltaSeconds 限制为 0.05；该取舍会在卡顿时减慢游戏时间"),
        ], [2.5, 4.2, 6.2, 6.3], ["center", "left", "left", "left"], 8.3)
    add_body(doc, "当前实现将输入、状态更新和绘制分开，主要代码如下。")
    add_code(doc, "double deltaSeconds = lastTime == 0\n        ? 0\n        : Math.min((now - lastTime) / 1_000_000_000.0, 0.05);\nlastTime = now;\ncontroller.update(deltaSeconds);\nrenderFrame(graphics);")

    add_heading(doc, "4.3 AI 核对总结", 2)
    add_table(doc,
        ["项目", "结论"],
        [
            ("AI 生成方法数", "不确定。AI 历史记录包含多段完整示例，但未标记哪些方法被直接采用"),
            ("本次可核对示例", "2 组：倒计时设计、Canvas 移动设计"),
            ("发现并已处理的问题", "可确认 4 类改进：职责耦合、双时间源、帧率相关移动、绘制与状态混合；当前源码已分别通过分层与 deltaSeconds 处理"),
            ("仍未解决的问题", "暂停恢复、订单与烹饪闭环、结果页 UI、单元测试、对角线速度归一化、参数空值或负值校验"),
            ("后续使用 AI 注意事项", "保留提示词和原始输出；逐方法记录采用与修改情况；合入前编译并补充单元测试；不把示例目标写成已实现事实"),
        ], [5.0, 13.2], ["left", "left"], 8.8)

    # Make body styles deterministic while preserving the template's visual family.
    for p in doc.paragraphs:
        if p.style and p.style.name in ("Heading 1", "Heading 2"):
            keep_with_next(p)
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    core = doc.core_properties
    core.title = "出餐 出餐 详细设计说明书"
    core.subject = "JavaFX 游戏项目详细设计"
    core.keywords = "JavaFX, 详细设计, 状态机, 方法设计"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    if sha256(REFERENCE) != EXPECTED_SHA256:
        raise RuntimeError("Reference template changed during authoring.")
    print(OUTPUT)


if __name__ == "__main__":
    build()
